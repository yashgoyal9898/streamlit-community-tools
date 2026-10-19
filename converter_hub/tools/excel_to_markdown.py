"""Excel (.xlsx) -> Markdown.

Hybrid approach:
  * pandas   - detects the real used range (no phantom empty rows / columns)
  * openpyxl - reads cells with their formatting (%, dates, currency, ...)
"""
from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from converter_hub.core.models import ConversionResult, Tool

# ═══════════════════════════════════════════════════════════════════════
#  Generic helpers
# ═══════════════════════════════════════════════════════════════════════


def _escape(value) -> str:
    """Make a value safe for a Markdown table cell."""
    if value is None:
        return ""
    text = str(value).replace("|", "\\|")
    return text.replace("\r\n", "<br>").replace("\n", "<br>")


def _positive_fmt(fmt: str) -> str:
    """Excel formats look like 'positive;negative;zero'. Keep the first part."""
    return fmt.split(";")[0] if ";" in fmt else fmt


def _decimal_places(fmt: str) -> int:
    m = re.search(r"\.([0#]+)", fmt)
    return len(m.group(1)) if m else 0


# ═══════════════════════════════════════════════════════════════════════
#  Date formatting  (single-pass tokenizer)
# ═══════════════════════════════════════════════════════════════════════

_DATE_TOKEN_RE = re.compile(
    r'"(?P<lit>[^"]*)"'                       # "quoted literal"
    r"|\\(?P<esc>.)"                          # \escaped char
    r"|(?P<tok>AM/PM|A/P|dddd|ddd|dd|d|yyyy|yy|mmmm|mmm|mm|m|hh|h|ss|s)"
    r"|(?P<other>.)",
    re.IGNORECASE,
)

_STRFTIME = {
    "dddd": "%A", "ddd": "%a", "dd": "%d", "d": "%d",
    "yyyy": "%Y", "yy": "%y",
    "mmmm": "%B", "mmm": "%b",
    "ss": "%S", "s": "%S",
    "am/pm": "%p", "a/p": "%p",
}


def _excel_date_fmt_to_strftime(fmt: str) -> str:
    """Excel date format -> Python strftime format.

    Handles the Excel quirk that 'm'/'mm' means *minutes* when it sits next to
    hours or seconds, and *month* otherwise. Uses 12-hour clock if AM/PM exists.
    """
    fmt = re.sub(r"\[\$[^\]]*\]", "", fmt)                       # [$-409] locale tags
    fmt = re.sub(r"\[(h+|m+|s+)\]", r"\1", fmt, flags=re.IGNORECASE)  # [h] elapsed

    parts: list[tuple[str, str]] = []
    for m in _DATE_TOKEN_RE.finditer(fmt):
        if m.group("lit") is not None:
            parts.append(("lit", m.group("lit")))
        elif m.group("esc") is not None:
            parts.append(("lit", m.group("esc")))
        elif m.group("tok"):
            parts.append(("tok", m.group("tok").lower()))
        else:
            parts.append(("lit", m.group("other")))

    token_idx = [i for i, (kind, _) in enumerate(parts) if kind == "tok"]
    twelve_hour = any(parts[i][1] in ("am/pm", "a/p") for i in token_idx)
    hour_code = "%I" if twelve_hour else "%H"

    out: list[str] = []
    for i, (kind, text) in enumerate(parts):
        if kind == "lit":
            out.append(text.replace("%", "%%"))
        elif text in ("h", "hh"):
            out.append(hour_code)
        elif text in ("m", "mm"):
            pos = token_idx.index(i)
            prev_tok = parts[token_idx[pos - 1]][1] if pos > 0 else ""
            next_tok = parts[token_idx[pos + 1]][1] if pos + 1 < len(token_idx) else ""
            is_minute = prev_tok in ("h", "hh") or next_tok in ("s", "ss")
            out.append("%M" if is_minute else "%m")
        else:
            out.append(_STRFTIME[text])
    return "".join(out)


def _format_date(value: date | datetime, fmt: str) -> str:
    try:
        return value.strftime(_excel_date_fmt_to_strftime(_positive_fmt(fmt)))
    except Exception:
        if isinstance(value, datetime) and (value.hour or value.minute or value.second):
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d")


# ═══════════════════════════════════════════════════════════════════════
#  Number formatting
# ═══════════════════════════════════════════════════════════════════════

_ESCAPE_CHARS_RE = re.compile(r'["\\_*]')
_PREFIX_RE = re.compile(r"^([^\d#,.0@?+\-\s]+)")


def _clean_number_fmt(fmt: str) -> str:
    fmt = re.sub(r"\[\$([^\]\-]*)[^\]]*\]", r"\1", fmt)  # [$₹-4009] -> ₹
    fmt = re.sub(r"\[[^\]]*\]", "", fmt)                 # [Red], [Blue] ...
    return _ESCAPE_CHARS_RE.sub("", fmt)


def _format_number(value: int | float, fmt: str) -> str:
    fmt = _positive_fmt(fmt)

    if fmt in ("General", "@", ""):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)

    use_comma = "#,##" in fmt
    decimals = _decimal_places(fmt)

    pm = _PREFIX_RE.match(_clean_number_fmt(fmt))
    prefix = pm.group(1).strip() if pm else ""

    try:
        if use_comma:
            formatted = f"{value:,.{decimals}f}"
        elif "#" in fmt or "0" in fmt:
            formatted = f"{value:.{decimals}f}"
        else:
            return str(value)
        return f"{prefix}{formatted}"
    except Exception:
        return str(value)


# ═══════════════════════════════════════════════════════════════════════
#  Cell / sheet formatting
# ═══════════════════════════════════════════════════════════════════════


def _format_cell(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return _escape(value)

    fmt = str(cell.number_format or "General")

    if "%" in fmt and isinstance(value, (int, float)):
        decimals = _decimal_places(_positive_fmt(fmt))
        return _escape(f"{value * 100:.{decimals}f}%")

    if isinstance(value, (datetime, date)):
        return _escape(_format_date(value, fmt))

    if isinstance(value, (int, float)):
        return _escape(_format_number(value, fmt))

    return _escape(value)


def _get_data_bounds(excel_bytes: BytesIO) -> dict[str, tuple[int, int]]:
    """Real (rows, cols) of every sheet.

    openpyxl's max_row / max_column also count formatting-only cells, which can
    create thousands of empty Markdown columns. pandas ignores trailing empties.
    """
    excel_bytes.seek(0)
    xl = pd.ExcelFile(excel_bytes, engine="openpyxl")
    bounds: dict[str, tuple[int, int]] = {}

    for sheet_name in xl.sheet_names:
        try:
            df = xl.parse(sheet_name, header=None, dtype=object)
        except Exception:
            bounds[sheet_name] = (0, 0)
            continue

        if df.empty:
            bounds[sheet_name] = (0, 0)
            continue

        has_data = df.notna()
        row_mask = has_data.any(axis=1)
        col_mask = has_data.any(axis=0)

        if not row_mask.any():
            bounds[sheet_name] = (0, 0)
            continue

        last_row = int(row_mask[::-1].idxmax()) + 1
        last_col = int(col_mask[::-1].idxmax()) + 1
        bounds[sheet_name] = (last_row, last_col)

    return bounds


def _worksheet_to_markdown(ws: Worksheet, nrows: int, ncols: int) -> str:
    if nrows == 0 or ncols == 0:
        return "_Empty Sheet_"

    rows = list(ws.iter_rows(min_row=1, max_row=nrows, min_col=1, max_col=ncols))
    if not rows:
        return "_Empty Sheet_"

    header = [_format_cell(c) for c in rows[0]]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    for row in rows[1:]:
        lines.append("| " + " | ".join(_format_cell(c) for c in row) + " |")
    return "\n".join(lines)


# ═══════════════════════════════════════════════════════════════════════
#  Public API + tool registration
# ═══════════════════════════════════════════════════════════════════════


def excel_bytes_to_markdown(excel_bytes: BytesIO, workbook_name: str) -> str:
    bounds = _get_data_bounds(excel_bytes)

    excel_bytes.seek(0)
    wb = load_workbook(excel_bytes, data_only=True)  # data_only -> formula results

    md: list[str] = [f"# {workbook_name}", ""]
    for ws in wb.worksheets:
        nrows, ncols = bounds.get(ws.title, (0, 0))
        md.extend([f"## {ws.title}", "", _worksheet_to_markdown(ws, nrows, ncols), "", "---", ""])
    return "\n".join(md)


def convert(data: bytes, filename: str) -> ConversionResult:
    stem = Path(filename).stem
    markdown = excel_bytes_to_markdown(BytesIO(data), stem)
    return ConversionResult(
        data=markdown,
        filename=f"{stem}.md",
        mime="text/markdown",
        preview=markdown,
    )


TOOL = Tool(
    slug="excel-to-markdown",
    title="Excel → Markdown",
    icon="📊",
    description="Convert every sheet of an .xlsx workbook into Markdown tables.",
    file_types=("xlsx",),
    convert=convert,
    order=20,
)
