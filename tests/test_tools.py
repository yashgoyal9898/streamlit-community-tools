import zipfile
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from converter_hub.core.registry import discover_tools
from converter_hub.tools import csv_to_excel, excel_to_markdown, zip_to_markdown


def test_registry_discovers_all_tools():
    slugs = [t.slug for t in discover_tools()]
    assert slugs == ["zip-to-markdown", "excel-to-markdown", "csv-to-excel"]


def test_zip_to_markdown():
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a/main.py", "print('hi')\n")
        z.writestr("README.md", "# T\n```py\nx=1\n```\n")   # contains fences
        z.writestr("img.bin", b"\x00\x01\x02")               # binary
        z.writestr("__MACOSX/junk", "x")                     # ignored
    md = zip_to_markdown.convert(buf.getvalue(), "proj.zip").data
    assert "## 📄 a/main.py" in md and "```py\nprint('hi')" in md
    assert "````md" in md                      # longer fence for nested ```
    assert "Skipped: binary file" in md
    assert "__MACOSX" not in md


def test_zip_invalid():
    with pytest.raises(ValueError):
        zip_to_markdown.convert(b"not a zip", "x.zip")


def test_excel_to_markdown():
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Pct", "Price", "When", "Time"])
    ws.append(["A|B", 0.256, 1234.5, datetime(2026, 10, 19), datetime(2026, 1, 1, 14, 5, 9)])
    ws["B2"].number_format = "0.0%"
    ws["C2"].number_format = '"₹"#,##0.00'
    ws["D2"].number_format = "dd-mmm-yyyy"
    ws["E2"].number_format = "hh:mm:ss AM/PM"
    ws["Z100"].number_format = "0.00"        # formatting-only cell must not widen table
    buf = BytesIO()
    wb.save(buf)

    md = excel_to_markdown.convert(buf.getvalue(), "book.xlsx").data
    assert "| Name | Pct | Price | When | Time |" in md
    assert "A\\|B | 25.6% | ₹1,234.50 | 19-Oct-2026 | 02:05:09 PM |" in md
    assert md.count("---") >= 6  # separator row + sheet divider


def test_excel_leading_blank_rows_and_cols():
    wb = Workbook()
    ws = wb.active
    ws["C3"] = "x"
    ws["D4"] = "y"
    buf = BytesIO()
    wb.save(buf)
    md = excel_to_markdown.convert(buf.getvalue(), "b.xlsx").data
    assert "x" in md and "y" in md


def test_csv_to_excel_utf8_and_latin1():
    out = csv_to_excel.convert("a,b\n1,é\n".encode("utf-8"), "d.csv")
    assert out.filename == "d.xlsx"
    assert load_workbook(BytesIO(out.data)).active["B2"].value == "é"

    out = csv_to_excel.convert("a,b\n1,é\n".encode("latin-1"), "d.csv")
    assert load_workbook(BytesIO(out.data)).active["B2"].value == "é"
